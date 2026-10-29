"""
Exercício 14.07 — Ler a planilha de valores (4 abas) com openpyxl
Nível: ★★★

Contexto: O laboratório preenche um .xlsx com uma aba por linha de produção
("Linha A" ... "Linha D"), cabeçalho na linha 1 e uma OP por linha. Gente digita
cabeçalho de jeitos diferentes ("Espessura (µm)", "ESPESSURA (um)") e às vezes digita
número como texto ("41,2"). O programa deve ser tolerante no FORMATO e rígido no
CONTEÚDO.

Tarefa: implemente três funções.
1. `chave_coluna(cabecalho)`: apague o que estiver entre parênteses (unidade),
   normalize (peça pronta `normalizar`) e troque espaços por "_".
   "Espessura nominal (µm)" -> "espessura_nominal"; None -> "".
2. `numero_celula(valor)`: int ou float -> float; texto com vírgula ou ponto
   ("41,2", " 41.2 ") -> float; None, "" ou texto que não é número ("n/d") -> None.
   (bool é subclasse de int: True/False também devolvem None.)
3. `ler_planilha(caminho)` devolve list[Medicao], na ordem das abas e das linhas:
   - abra com load_workbook(caminho, read_only=True, data_only=True) e FECHE no
     final (try/finally com .close());
   - leia só as abas cujo nome normalizado começa com "linha";
   - a 1ª linha é o cabeçalho: monte {chave_coluna(nome): índice da coluna};
   - se faltar alguma chave de OBRIGATORIAS, levante PlanilhaInvalida com uma
     mensagem que cite a aba e as colunas faltando;
   - pule linhas sem OP; para as demais crie Medicao com: op e linha via
     texto_celula; aba = nome da aba; linha_excel = número da linha no Excel
     (o cabeçalho é a 1, então a 1ª OP é a 2); valores = {característica: número}
     só das CARACTERISTICAS preenchidas; nominais = {"espessura": ..., "gramatura": ...}
     (colunas de NOMINAIS) só se preenchidos; especificacao, observacao e produto
     via texto_celula ("" quando vazio ou quando a coluna não existe).

Exemplos:
    >>> chave_coluna("Tração (MPa)")
    'tracao'
    >>> numero_celula("41,2")
    41.2
    >>> medicoes = ler_planilha(Path("dados/planilha_valores.xlsx"))
    >>> len(medicoes), medicoes[0].op, medicoes[0].aba, medicoes[0].linha_excel
    (16, '24518', 'Linha A', 2)

Dica: aba.iter_rows(values_only=True) devolve tuplas de valores; chame next() uma vez
para tirar o cabeçalho e use enumerate(..., start=2) no resto. Montar um dict por
linha ({chave: valor}) deixa o resto do código bem legível.

Teste com:  python curso.py testar 14 ex07
"""

import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

# ============================================================================
# PEÇAS PRONTAS (etapas anteriores) — não precisa mexer
# ============================================================================

CARACTERISTICAS = ("espessura", "gramatura", "cof", "tracao", "selagem")
NOMINAIS = {"espessura": "espessura_nominal", "gramatura": "gramatura_nominal"}
OBRIGATORIAS = {
    "op",
    "linha",
    "especificacao",
    "observacao",
    "espessura_nominal",
    "gramatura_nominal",
    *CARACTERISTICAS,
}


class PlanilhaInvalida(ValueError):
    """A planilha não tem o formato combinado (ex.: falta uma coluna)."""


@dataclass
class Medicao:
    op: str
    aba: str
    linha: str
    valores: dict[str, float] = field(default_factory=dict)
    nominais: dict[str, float] = field(default_factory=dict)
    especificacao: str = ""
    observacao: str = ""
    produto: str = ""
    linha_excel: int = 0


def normalizar(texto: str | None) -> str:  # etapa 6
    if texto is None:
        return ""
    decomposto = unicodedata.normalize("NFKD", str(texto))
    sem_acento = "".join(c for c in decomposto if not unicodedata.combining(c))
    return " ".join(sem_acento.lower().split())


def texto_celula(valor: object) -> str:
    """Texto limpo de uma célula. 24518.0 vira "24518"; None vira ""."""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip()


# ============================================================================
# SUA PARTE
# ============================================================================


def chave_coluna(cabecalho: object) -> str:
    """ "Espessura nominal (µm)" -> "espessura_nominal"."""
    # Tira a unidade ANTES de normalizar: assim "(µm)", "(um)" e "(micra)" somem
    # todos. O ".*?" é "preguiçoso": para no primeiro ")".
    sem_unidade = re.sub(r"\(.*?\)", " ", str(cabecalho or ""))
    return normalizar(sem_unidade).replace(" ", "_")


def numero_celula(valor: object) -> float | None:
    """Número da célula, aceitando texto com vírgula. Vazio ou inválido -> None."""
    # bool vem primeiro porque isinstance(True, int) é True em Python.
    if valor is None or isinstance(valor, bool):
        return None
    if isinstance(valor, (int, float)):
        return float(valor)
    texto = str(valor).strip().replace(",", ".")
    if not texto:
        return None
    try:
        return float(texto)
    except ValueError:
        # "n/d", "-", "refazer": não é número. None faz a validação dizer
        # "sem valor medido", que é exatamente o que aconteceu.
        return None


def ler_planilha(caminho: Path) -> list[Medicao]:
    """Lê as abas "Linha ..." e devolve uma Medicao por OP preenchida."""
    # read_only: lê sem carregar tudo na memória. data_only: traz o VALOR das
    # fórmulas (o que você vê no Excel), não o texto "=B2*2".
    pasta_trabalho = load_workbook(caminho, read_only=True, data_only=True)
    medicoes: list[Medicao] = []
    try:
        for aba in pasta_trabalho.worksheets:
            if not normalizar(aba.title).startswith("linha"):
                continue  # "Instruções", "Resumo"...: não são de produção
            linhas = aba.iter_rows(values_only=True)
            cabecalho = next(linhas, None) or ()
            colunas = {chave_coluna(nome): indice for indice, nome in enumerate(cabecalho)}
            faltando = sorted(OBRIGATORIAS - colunas.keys())  # diferença de conjuntos
            if faltando:
                raise PlanilhaInvalida(
                    f"aba {aba.title!r}: faltam colunas: {', '.join(faltando)}"
                )
            for numero, linha in enumerate(linhas, start=2):
                # Um dict por linha: registro["cof"] em vez de linha[9]. O "if" protege
                # contra linhas mais curtas que o cabeçalho.
                registro = {chave: linha[i] for chave, i in colunas.items() if i < len(linha)}
                op = texto_celula(registro.get("op"))
                if not op:
                    continue  # linha em branco (comum no fim da aba)
                valores = {}
                for caracteristica in CARACTERISTICAS:
                    valor = numero_celula(registro.get(caracteristica))
                    if valor is not None:
                        valores[caracteristica] = valor
                nominais = {}
                for caracteristica, coluna in NOMINAIS.items():
                    nominal = numero_celula(registro.get(coluna))
                    if nominal is not None:
                        nominais[caracteristica] = nominal
                medicoes.append(
                    Medicao(
                        op=op,
                        aba=aba.title,
                        linha=texto_celula(registro.get("linha")),
                        valores=valores,
                        nominais=nominais,
                        especificacao=texto_celula(registro.get("especificacao")),
                        observacao=texto_celula(registro.get("observacao")),
                        produto=texto_celula(registro.get("produto")),
                        linha_excel=numero,
                    )
                )
    finally:
        # Em modo read_only o arquivo fica aberto até o close(). No Windows, isso
        # impede salvar a planilha no Excel enquanto o programa roda.
        pasta_trabalho.close()
    return medicoes
