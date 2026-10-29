"""Espiar a planilha de valores antes de escrever o leitor de verdade.

Primeiro passo de qualquer integração: OLHAR o arquivo real com o Python.
Quais abas existem? Como está o cabeçalho? Que tipo cada célula traz?

Rode:  python exemplos/04_espiar_planilha.py
"""

from pathlib import Path

from openpyxl import load_workbook

PLANILHA = Path(__file__).resolve().parent.parent / "dados" / "planilha_valores.xlsx"

pasta_trabalho = load_workbook(PLANILHA, read_only=True, data_only=True)
try:
    print("abas:", pasta_trabalho.sheetnames)
    aba = pasta_trabalho["Linha A"]
    linhas = aba.iter_rows(values_only=True)
    cabecalho = next(linhas)
    print("cabeçalho:", cabecalho[:5], "...")
    primeira = next(linhas)
    for nome, valor in list(zip(cabecalho, primeira, strict=True))[:7]:
        print(f"  {nome!r:28} -> {valor!r} ({type(valor).__name__})")
    for aba in pasta_trabalho.worksheets:
        ops = [linha[0] for linha in aba.iter_rows(min_row=2, values_only=True) if linha[0]]
        print(f"{aba.title}: {len(ops)} OPs {ops}")
finally:
    pasta_trabalho.close()
