"""
Exercício 14.07 — Ler a planilha de valores (4 abas) com openpyxl
Nível: ★★★

Contexto: O laboratório preenche um .xlsx com uma aba por linha de produção
("Linha A" ... "Linha D"), cabeçalho na linha 1 e uma OP por linha. Gente digita
cabeçalho de jeitos diferentes ("Espessura (µm)", "ESPESSURA (um)") e às vezes digita
número como texto ("41,2"). O programa deve ser tolerante no FORMATO e rígido no
CONTEÚDO.

Tarefa: implemente três funções.
1. `chave_coluna(cabecalho)`: apague o que estiver entre parênteses (unidade),
   normalize (peça pronta `normalizar`) e troque espaços por "_".
   "Espessura nominal (µm)" -> "espessura_nominal"; None -> "".
2. `numero_celula(valor)`: int ou float -> float; texto com vírgula ou ponto
   ("41,2", " 41.2 ") -> float; None, "" ou texto que não é número ("n/d") -> None.
   (bool é subclasse de int: True/False também devolvem None.)
3. `ler_planilha(caminho)` devolve list[Medicao], na ordem das abas e das linhas:
   - abra com load_workbook(caminho, read_only=True, data_only=True) e FECHE no
     final (try/finally com .close());
   - leia só as abas cujo nome normalizado começa com "linha";
   - a 1ª linha é o cabeçalho: monte {chave_coluna(nome): índice da coluna};
   - se faltar alguma chave de OBRIGATORIAS, levante PlanilhaInvalida com uma
     mensagem que cite a aba e as colunas faltando;
   - pule linhas sem OP; para as demais crie Medicao com: op e linha via
     texto_celula; aba = nome da aba; linha_excel = número da linha no Excel
     (o cabeçalho é a 1, então a 1ª OP é a 2); valores = {característica: número}
     só das CARACTERISTICAS preenchidas; nominais = {"espessura": ..., "gramatura": ...}
     (colunas de NOMINAIS) só se preenchidos; especificacao, observacao e produto
     via texto_celula ("" quando vazio ou quando a coluna não existe).

Exemplos:
    >>> chave_coluna("Tração (MPa)")
    'tracao'
    >>> numero_celula("41,2")
    41.2
    >>> medicoes = ler_planilha(Path("dados/planilha_valores.xlsx"))
    >>> len(medicoes), medicoes[0].op, medicoes[0].aba, medicoes[0].linha_excel
    (16, '24518', 'Linha A', 2)

Dica: aba.iter_rows(values_only=True) devolve tuplas de valores; chame next() uma vez
para tirar o cabeçalho e use enumerate(..., start=2) no resto. Montar um dict por
linha ({chave: valor}) deixa o resto do código bem legível.

Teste com:  python curso.py testar 14 ex07
"""

import re  # noqa: F401 (você vai usar)
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook  # noqa: F401 (você vai usar)

# ============================================================================
# PEÇAS PRONTAS (etapas anteriores) — não precisa mexer
# ============================================================================

CARACTERISTICAS = ("espessura", "gramatura", "cof", "tracao", "selagem")
NOMINAIS = {"espessura": "espessura_nominal", "gramatura": "gramatura_nominal"}
OBRIGATORIAS = {
    "op",
    "linha",
    "especificacao",
    "observacao",
    "espessura_nominal",
    "gramatura_nominal",
    *CARACTERISTICAS,
}


class PlanilhaInvalida(ValueError):
    """A planilha não tem o formato combinado (ex.: falta uma coluna)."""


@dataclass
class Medicao:
    op: str
    aba: str
    linha: str
    valores: dict[str, float] = field(default_factory=dict)
    nominais: dict[str, float] = field(default_factory=dict)
    especificacao: str = ""
    observacao: str = ""
    produto: str = ""
    linha_excel: int = 0


def normalizar(texto: str | None) -> str:  # etapa 6
    if texto is None:
        return ""
    decomposto = unicodedata.normalize("NFKD", str(texto))
    sem_acento = "".join(c for c in decomposto if not unicodedata.combining(c))
    return " ".join(sem_acento.lower().split())


def texto_celula(valor: object) -> str:
    """Texto limpo de uma célula. 24518.0 vira "24518"; None vira ""."""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip()


# ============================================================================
# SUA PARTE
# ============================================================================


def chave_coluna(cabecalho: object) -> str:
    """ "Espessura nominal (µm)" -> "espessura_nominal"."""
    # Escreva seu código aqui
    raise NotImplementedError("Implemente esta função")


def numero_celula(valor: object) -> float | None:
    """Número da célula, aceitando texto com vírgula. Vazio ou inválido -> None."""
    # Escreva seu código aqui
    raise NotImplementedError("Implemente esta função")


def ler_planilha(caminho: Path) -> list[Medicao]:
    """Lê as abas "Linha ..." e devolve uma Medicao por OP preenchida."""
    # Escreva seu código aqui
    raise NotImplementedError("Implemente esta função")
