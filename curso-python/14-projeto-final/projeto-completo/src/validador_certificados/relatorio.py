"""Relatório da série: CSV (simples, bom para sistemas) e Excel (para gente)."""

from __future__ import annotations

import csv
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill

from .modelos import ItemFila

COLUNAS = ["OP", "NF", "Peso (kg)", "Especificação", "Situação", "Motivos", "Comprovante"]

ROTULOS = {
    "emitir": "A EMITIR",
    "bloqueado": "BLOQUEADO",
    "ensaio": "ENSAIO (não emitido)",
    "emitido": "EMITIDO",
    "falha": "FALHA",
    "ignorado": "JÁ EMITIDO",
}

CORES = {  # fundo da linha no Excel, por situação
    "bloqueado": "F8D7DA",
    "falha": "F8D7DA",
    "ensaio": "D1E7DD",
    "emitido": "D1E7DD",
    "ignorado": "E2E3E5",
}


def resumir(itens: list[ItemFila]) -> dict[str, int]:
    """Quantos itens em cada situação: {"bloqueado": 4, "ensaio": 2}."""
    return dict(sorted(Counter(item.situacao for item in itens).items()))


def escrever_relatorio_csv(itens: list[ItemFila], caminho: Path) -> Path:
    """CSV com ';' e utf-8-sig: abre direto no Excel em português."""
    caminho = Path(caminho)
    caminho.parent.mkdir(parents=True, exist_ok=True)
    with caminho.open("w", encoding="utf-8-sig", newline="") as arquivo:
        escritor = csv.writer(arquivo, delimiter=";")
        escritor.writerow(COLUNAS)
        for item in itens:
            escritor.writerow(
                [
                    item.op,
                    item.nf,
                    f"{item.peso_kg:.2f}".replace(".", ","),
                    item.especificacao,
                    item.situacao.upper(),
                    " | ".join(item.motivos),
                    item.comprovante,
                ]
            )
    return caminho


def gerar_relatorio_excel(
    itens: list[ItemFila], fora_do_padrao: dict[str, str], caminho: Path
) -> Path:
    """Excel com 3 abas: Série, E-mails fora do padrão e Resumo (com gráfico)."""
    caminho = Path(caminho)
    caminho.parent.mkdir(parents=True, exist_ok=True)
    pasta_trabalho = Workbook()

    serie = pasta_trabalho.active
    serie.title = "Série"
    serie.append(COLUNAS)
    for item in itens:
        serie.append(
            [
                item.op,
                item.nf,
                item.peso_kg,
                item.especificacao,
                ROTULOS.get(item.situacao, item.situacao),
                "\n".join(item.motivos),
                item.comprovante,
            ]
        )
        cor = CORES.get(item.situacao)
        if cor:
            for celula in serie[serie.max_row]:
                celula.fill = PatternFill("solid", fgColor=cor)
        serie.cell(serie.max_row, 3).number_format = "#,##0.00"
    _formatar(serie, [10, 12, 12, 16, 22, 60, 22])

    emails = pasta_trabalho.create_sheet("E-mails fora do padrão")
    emails.append(["Arquivo", "Problema"])
    for arquivo, problema in fora_do_padrao.items():
        emails.append([arquivo, problema])
    _formatar(emails, [28, 70])

    resumo = pasta_trabalho.create_sheet("Resumo")
    resumo.append(["Situação", "Quantidade"])
    for situacao, quantidade in resumir(itens).items():
        resumo.append([ROTULOS.get(situacao, situacao), quantidade])
    resumo.append(["E-MAIL FORA DO PADRÃO", len(fora_do_padrao)])
    _formatar(resumo, [26, 12])
    _grafico_resumo(resumo)

    pasta_trabalho.save(caminho)
    return caminho


def _formatar(aba, larguras: list[int]) -> None:
    for celula in aba[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="0D4A6B")
    for indice, largura in enumerate(larguras, start=1):
        aba.column_dimensions[aba.cell(1, indice).column_letter].width = largura
    aba.freeze_panes = "A2"


def _grafico_resumo(aba) -> None:
    """Gráfico de barras nativo do Excel (continua editável lá dentro)."""
    if aba.max_row < 2:
        return
    grafico = BarChart()
    grafico.title = "Situação da série"
    grafico.y_axis.title = "Quantidade"
    grafico.legend = None
    dados = Reference(aba, min_col=2, min_row=1, max_row=aba.max_row)
    categorias = Reference(aba, min_col=1, min_row=2, max_row=aba.max_row)
    grafico.add_data(dados, titles_from_data=True)
    grafico.set_categories(categorias)
    aba.add_chart(grafico, "D2")
