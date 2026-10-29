"""Leitura da planilha de valores (openpyxl) e conferência da linha de produção.

Regra de ouro desta camada: ser tolerante no FORMATO (maiúsculas, acentos, unidade no
cabeçalho, número digitado como texto) e rígido no CONTEÚDO (coluna faltando ou linha
divergente nunca é "adivinhada").
"""

from __future__ import annotations

import logging
import re
from pathlib import Path

from openpyxl import load_workbook

from .modelos import Medicao, PlanilhaInvalida
from .texto import normalizar

logger = logging.getLogger(__name__)

LINHAS_VALIDAS = "ABCD"
CARACTERISTICAS = ("espessura", "gramatura", "cof", "tracao", "selagem")
NOMINAIS = {"espessura": "espessura_nominal", "gramatura": "gramatura_nominal"}
OBRIGATORIAS = {
    "op",
    "linha",
    "especificacao",
    "observacao",
    *CARACTERISTICAS,
    *NOMINAIS.values(),
}


# --------------------------------------------------------------------------- células
def chave_coluna(cabecalho: object) -> str:
    """Chave da coluna: "Espessura nominal (µm)" -> "espessura_nominal"."""
    sem_unidade = re.sub(r"\(.*?\)", " ", str(cabecalho or ""))
    return normalizar(sem_unidade).replace(" ", "_")


def texto_celula(valor: object) -> str:
    """Texto limpo de uma célula. 24518.0 vira "24518"; None vira ""."""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip()


def numero_celula(valor: object) -> float | None:
    """Número de uma célula. Aceita texto com vírgula ("41,2"). Vazio ou inválido -> None."""
    if valor is None or isinstance(valor, bool):
        return None
    if isinstance(valor, (int, float)):
        return float(valor)
    texto = str(valor).strip().replace(",", ".")
    if not texto:
        return None
    try:
        return float(texto)
    except ValueError:
        return None


# --------------------------------------------------------------------------- leitura
def ler_planilha(caminho: Path) -> list[Medicao]:
    """Lê as abas "Linha ..." e devolve uma Medicao por OP preenchida."""
    caminho = Path(caminho)
    if not caminho.is_file():
        raise FileNotFoundError(f"planilha não encontrada: {caminho}")
    pasta_trabalho = load_workbook(caminho, read_only=True, data_only=True)
    medicoes: list[Medicao] = []
    try:
        for aba in pasta_trabalho.worksheets:
            if not normalizar(aba.title).startswith("linha"):
                logger.debug("aba %r ignorada (não é de linha de produção)", aba.title)
                continue
            medicoes.extend(_ler_aba(aba))
    finally:
        pasta_trabalho.close()  # em read_only o arquivo fica aberto até fechar
    logger.info("planilha: %d OPs lidas de %s", len(medicoes), caminho)
    return medicoes


def _ler_aba(aba) -> list[Medicao]:
    linhas = aba.iter_rows(values_only=True)
    cabecalho = next(linhas, None) or ()
    colunas = {chave_coluna(nome): indice for indice, nome in enumerate(cabecalho)}
    faltando = sorted(OBRIGATORIAS - colunas.keys())
    if faltando:
        raise PlanilhaInvalida(f"aba {aba.title!r}: faltam colunas: {', '.join(faltando)}")

    medicoes = []
    for numero, linha in enumerate(linhas, start=2):
        # {"op": 24518, "linha": "A", "cof": 0.22, ...} -- só as colunas que existem
        registro = {chave: linha[i] for chave, i in colunas.items() if i < len(linha)}
        op = texto_celula(registro.get("op"))
        if not op:
            continue  # linha vazia (comum no fim da aba)
        valores = {}
        for caracteristica in CARACTERISTICAS:
            valor = numero_celula(registro.get(caracteristica))
            if valor is not None:
                valores[caracteristica] = valor
        nominais = {}
        for caracteristica, coluna in NOMINAIS.items():
            nominal = numero_celula(registro.get(coluna))
            if nominal is not None:
                nominais[caracteristica] = nominal
        medicoes.append(
            Medicao(
                op=op,
                aba=aba.title,
                linha=texto_celula(registro.get("linha")),
                valores=valores,
                nominais=nominais,
                especificacao=texto_celula(registro.get("especificacao")),
                observacao=texto_celula(registro.get("observacao")),
                produto=texto_celula(registro.get("produto")),
                linha_excel=numero,
            )
        )
    return medicoes


# --------------------------------------------------------------------------- linha
def normalizar_linha(texto: object) -> str | None:
    """Letra da linha: "Linha A", "linha  b", " c " -> "A", "B", "C". Outra coisa -> None."""
    if texto is None:
        return None
    achado = re.fullmatch(r"(?:LINHA\s*)?([A-Z])", str(texto).strip().upper())
    if achado and achado.group(1) in LINHAS_VALIDAS:
        return achado.group(1)
    return None


def conferir_linha(nome_aba: str, linha_informada: object) -> str | None:
    """Confere se a linha digitada na planilha bate com a aba. None = tudo certo."""
    da_aba = normalizar_linha(nome_aba)
    if da_aba is None:
        return f"aba {nome_aba!r} não corresponde a uma linha de produção"
    informada = normalizar_linha(linha_informada)
    if informada is None:
        texto = "" if linha_informada is None else str(linha_informada)
        return f"linha não informada ou inválida: {texto!r}"
    if informada != da_aba:
        return f"linha informada {informada} diferente da aba {nome_aba}"
    return None
