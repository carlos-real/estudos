import csv

from openpyxl import load_workbook

from validador_certificados.historico import carregar_emitidas, registrar_emitidas
from validador_certificados.modelos import ItemFila
from validador_certificados.relatorio import (
    escrever_relatorio_csv,
    gerar_relatorio_excel,
    resumir,
)

ITENS = [
    ItemFila("24518", "123456", 1250.5, "ensaio", "PADRAO-A", comprovante="ENSAIO-1"),
    ItemFila("24519", "123457", 980.0, "bloqueado", "PADRAO-A", ["espessura alta", "COF"]),
]


def test_resumir():
    assert resumir(ITENS) == {"bloqueado": 1, "ensaio": 1}
    assert resumir([]) == {}


def test_csv_no_formato_do_excel_brasileiro(tmp_path):
    caminho = escrever_relatorio_csv(ITENS, tmp_path / "sub" / "serie.csv")
    with caminho.open(encoding="utf-8-sig", newline="") as arquivo:
        linhas = list(csv.reader(arquivo, delimiter=";"))
    assert linhas[0][0] == "OP"
    assert linhas[1] == ["24518", "123456", "1250,50", "PADRAO-A", "ENSAIO", "", "ENSAIO-1"]
    assert linhas[2][5] == "espessura alta | COF"


def test_excel_tem_tres_abas(tmp_path):
    caminho = gerar_relatorio_excel(ITENS, {"x.txt": "OP não encontrada"}, tmp_path / "r.xlsx")
    pasta_trabalho = load_workbook(caminho)
    assert pasta_trabalho.sheetnames == ["Série", "E-mails fora do padrão", "Resumo"]
    serie = pasta_trabalho["Série"]
    assert serie["E2"].value == "ENSAIO (não emitido)"
    assert serie["F3"].value == "espessura alta\nCOF"
    assert pasta_trabalho["E-mails fora do padrão"]["B2"].value == "OP não encontrada"


def test_historico(tmp_path):
    caminho = tmp_path / "emitidos.csv"
    assert carregar_emitidas(caminho) == set()
    itens = [
        ItemFila("1", "10", 5.0, "emitido", comprovante="CQ-1"),
        ItemFila("2", "20", 5.0, "falha"),
    ]
    assert registrar_emitidas(caminho, itens) == 1
    assert registrar_emitidas(caminho, itens) == 1  # acrescenta sem repetir cabeçalho
    assert carregar_emitidas(caminho) == {("1", "10")}
    assert caminho.read_text(encoding="utf-8").count("op;nf") == 1
