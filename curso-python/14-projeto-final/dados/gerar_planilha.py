"""Gera a planilha de valores de exemplo: dados/planilha_valores.xlsx.

Uso (na pasta curso-python):
    python 14-projeto-final/dados/gerar_planilha.py

São 4 abas (Linha A, B, C, D), uma linha por OP, no formato que o laboratório
preenche. Os dados são fictícios, mas cobrem os casos que o sistema precisa tratar:
OPs aprovadas, reprovadas, sem liberação de COF, com especificação de cliente,
com especificação desconhecida, com linha divergente da aba e com valor faltando.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

PASTA = Path(__file__).resolve().parent
DESTINO = PASTA / "planilha_valores.xlsx"

CABECALHO = [
    "OP",
    "Data",
    "Linha",
    "Produto",
    "Especificação",
    "Espessura nominal (µm)",
    "Espessura (µm)",
    "Gramatura nominal (g/m²)",
    "Gramatura (g/m²)",
    "COF",
    "Tração (MPa)",
    "Selagem (N/15mm)",
    "Observação",
]

D = date  # só para encurtar as linhas abaixo

# OP, data, linha, produto, especificação, esp. nominal, esp., gram. nominal, gram.,
# COF, tração, selagem, observação
ABAS = {
    "Linha A": [
        (24518, D(2026, 10, 20), "A", "Filme PE 40 µm transparente", "",
         40, 41.2, 37.0, 37.9, 0.22, 24.5, 9.8, "COF: liberado"),
        (24519, D(2026, 10, 20), "A", "Filme PE 40 µm transparente", "",
         40, 44.8, 37.0, 38.1, 0.24, 23.9, 9.1, "COF: liberado"),
        (24520, D(2026, 10, 21), "A", "Filme PE 50 µm branco", "",
         50, 49.1, 46.3, 45.0, 0.31, 22.0, 10.2, "cof liberado"),
        (24521, D(2026, 10, 21), "A", "Filme PE 60 µm sacaria", "ESP-0533",
         60, 61.0, 55.5, 56.0, 0.28, 29.0, 11.0, "COF: Liberado"),
        (24522, D(2026, 10, 22), "A", "Filme PE 40 µm transparente", "",
         40, 39.5, 37.0, 36.8, 0.38, 21.5, 7.6, "COF: liberado"),
    ],
    "Linha B": [
        (24527, D(2026, 10, 21), "B", "Filme PE 35 µm", "",
         35, 35.4, 32.4, 32.0, 0.30, 19.5, 6.8, "COF: aguardando 24h"),
        (24528, D(2026, 10, 21), "B", "Filme PE 35 µm", "",
         35, 34.0, 32.4, 33.1, 0.27, 20.1, 7.2, "COF:LIBERADO"),
        (24533, D(2026, 10, 22), "C", "Filme PE 45 µm", "",
         45, 45.9, 41.6, 42.0, 0.25, 21.0, 7.0, "COF: liberado"),
        (24534, D(2026, 10, 23), "B", "Filme PE 30 µm", "",
         30, 30.6, 27.8, 28.0, 0.33, 18.7, 6.1, "COF: não liberado - refazer ensaio"),
    ],
    "Linha C": [
        (24541, D(2026, 10, 22), "C", "Coex 70 µm café 250 g", "ESP-0412",
         70, 71.5, 65.8, 66.4, 0.21, 27.3, 13.4, "cof liberado"),
        (24542, D(2026, 10, 22), "C", "Coex 70 µm café 250 g", "ESP-0412",
         70, 69.0, 65.8, 65.0, 0.27, 26.8, 12.9, "COF: liberado"),
        (24546, D(2026, 10, 23), "C", "Coex 60 µm", "ESP-9999",
         60, 60.5, 56.4, 56.0, 0.20, 26.0, 11.0, "COF: liberado"),
        (24547, D(2026, 10, 24), "C", "Coex 60 µm biscoito", "ESP-0618",
         60, 59.2, 56.4, 57.1, 0.16, 25.6, 10.8, "Cof: Liberado. Bobina 2 com rebarba"),
    ],
    "Linha D": [
        (24555, D(2026, 10, 23), "D", "Laminado PET/PE 62 µm", "",
         62, 63.0, 70.5, 71.2, 0.19, 34.0, 14.5, "COF - LIBERADO"),
        (24556, D(2026, 10, 23), "D", "Laminado PET/PE 62 µm", "",
         62, 62.4, 70.5, 64.9, 0.22, 33.1, 13.0, "COF: liberado"),
        (24560, D(2026, 10, 24), "D", "Laminado BOPP/PE 55 µm", "",
         55, 55.8, 52.0, 52.6, 0.24, 31.5, None, ""),
    ],
}  # fmt: skip

LARGURAS = [9, 12, 7, 30, 15, 12, 11, 13, 11, 7, 10, 11, 38]
FORMATOS = {7: "0.0", 9: "0.0", 10: "0.00", 11: "0.0", 12: "0.0"}  # coluna -> formato


def gerar(destino: Path = DESTINO) -> Path:
    """Cria a planilha e devolve o caminho."""
    pasta_trabalho = Workbook()
    pasta_trabalho.remove(pasta_trabalho.active)  # começa sem a aba "Sheet"
    cabecalho_fonte = Font(bold=True, color="FFFFFF")
    cabecalho_fundo = PatternFill("solid", fgColor="0D4A6B")
    for nome_aba, linhas in ABAS.items():
        aba = pasta_trabalho.create_sheet(nome_aba)
        aba.append(CABECALHO)
        for celula in aba[1]:
            celula.font = cabecalho_fonte
            celula.fill = cabecalho_fundo
            celula.alignment = Alignment(wrap_text=True, vertical="center")
        for linha in linhas:
            aba.append(list(linha))
        for indice, largura in enumerate(LARGURAS, start=1):
            aba.column_dimensions[aba.cell(1, indice).column_letter].width = largura
        for linha in aba.iter_rows(min_row=2):
            linha[1].number_format = "DD/MM/YYYY"
            for coluna, formato in FORMATOS.items():
                linha[coluna - 1].number_format = formato
        aba.freeze_panes = "B2"
    pasta_trabalho.save(destino)
    return destino


if __name__ == "__main__":
    caminho = gerar()
    total = sum(len(linhas) for linhas in ABAS.values())
    print(f"Planilha gerada: {caminho.name} ({len(ABAS)} abas, {total} OPs)")
