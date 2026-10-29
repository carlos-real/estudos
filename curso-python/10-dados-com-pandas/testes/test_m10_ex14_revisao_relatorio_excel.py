import pandas as pd
import pytest
from openpyxl import load_workbook
from pandas.testing import assert_frame_equal

from _curso import carregar

ex = carregar(__file__, "ex14_revisao_relatorio_excel")


@pytest.fixture
def resultados():
    R = ex.Resultado
    return [
        R("024447", "024447-01", 61.8, "APROVADA"),
        R("024447", "024447-03", 67.5, "REPROVADA"),
        R("024466", "024466-04", 56.9, "REPROVADA"),
        R("024488", "024488-01", 59.5, "APROVADA"),
    ]


def gerar_planilha(resultados, pasta):
    caminho = ex.exportar_relatorio(resultados, pasta / "saida" / "relatorio.xlsx")
    return load_workbook(caminho)


def cor(celula):
    return celula.fill.fgColor.rgb if celula.fill.fill_type == "solid" else None


def test_resultados_para_dataframe(resultados):
    df = ex.resultados_para_dataframe(resultados[:2])
    esperado = pd.DataFrame(
        {
            "op": ["024447", "024447"],
            "bobina": ["024447-01", "024447-03"],
            "espessura_um": [61.8, 67.5],
            "status": ["APROVADA", "REPROVADA"],
        }
    )
    assert_frame_equal(df, esperado, check_dtype=False)


def test_dataframe_vazio_tem_as_colunas():
    df = ex.resultados_para_dataframe([])
    assert len(df) == 0
    assert list(df.columns) == ["op", "bobina", "espessura_um", "status"]


def test_cria_pastas_e_devolve_o_caminho(resultados, tmp_path):
    destino = tmp_path / "a" / "b" / "relatorio.xlsx"
    devolvido = ex.exportar_relatorio(resultados, destino)
    assert devolvido == destino
    assert destino.exists()


def test_lista_vazia_da_value_error(tmp_path):
    with pytest.raises(ValueError):
        ex.exportar_relatorio([], tmp_path / "vazio.xlsx")


def test_aba_e_dados(resultados, tmp_path):
    planilha = gerar_planilha(resultados, tmp_path)
    assert planilha.sheetnames == ["Resultados"]
    aba = planilha["Resultados"]
    assert [c.value for c in aba[1]] == ["op", "bobina", "espessura_um", "status"], (
        "grave sem o índice: index=False"
    )
    assert aba.max_row == 5
    assert aba["B3"].value == "024447-03"
    assert aba["C3"].value == pytest.approx(67.5)


def test_cabecalho_negrito_e_congelado(resultados, tmp_path):
    planilha = gerar_planilha(resultados, tmp_path)
    aba = planilha["Resultados"]
    assert all(celula.font.bold for celula in aba[1]), "cabeçalho em negrito"
    assert aba.freeze_panes == "A2"


def test_largura_das_colunas(resultados, tmp_path):
    planilha = gerar_planilha(resultados, tmp_path)
    aba = planilha["Resultados"]
    for letra, cabecalho in zip(
        "ABCD", ["op", "bobina", "espessura_um", "status"], strict=True
    ):
        largura = aba.column_dimensions[letra].width
        assert largura >= len(cabecalho) + 2, f"coluna {letra} estreita demais"


def test_reprovadas_pintadas_e_aprovadas_sem_cor(resultados, tmp_path):
    planilha = gerar_planilha(resultados, tmp_path)
    aba = planilha["Resultados"]
    for numero_linha in (3, 4):  # as duas REPROVADAS
        cores = {cor(celula) for celula in aba[numero_linha]}
        assert len(cores) == 1 and cores.pop().endswith("FFC7CE"), (
            f"todas as células da linha {numero_linha} devem ter fundo FFC7CE"
        )
    for numero_linha in (2, 5):  # as APROVADAS
        assert all(cor(celula) is None for celula in aba[numero_linha])
