"""Exemplo 12 — Exportar: CSV para o Excel brasileiro e Excel formatado com openpyxl.

Rode (na pasta curso-python):
    python 10-dados-com-pandas/exemplos/12_exportar_excel_formatado.py

Os arquivos vão para 10-dados-com-pandas/saida/ (pasta ignorada pelo Git).
"""

from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

PASTA_MODULO = Path(__file__).resolve().parent.parent
DADOS = PASTA_MODULO / "dados"
SAIDA = PASTA_MODULO / "saida"
SAIDA.mkdir(exist_ok=True)

df = pd.read_csv(
    DADOS / "medicoes.csv", sep=";", decimal=",", thousands=".", dtype={"op": str}
)
specs = pd.read_csv(DADOS / "especificacoes.csv", sep=";", decimal=",")
df = df.merge(specs[["produto", "cof_max"]], on="produto", how="left")
df["status"] = "APROVADA"
df.loc[df["cof"] > df["cof_max"], "status"] = "REPROVADA"
liberada = df["observacao"].fillna("").str.contains("liberado", case=False)
df.loc[(df["status"] == "REPROVADA") & liberada, "status"] = "LIBERADA"
df.loc[df["cof_max"].isna(), "status"] = "SEM ESPECIFICAÇÃO"  # nunca aprove sem regra!
print(df["status"].value_counts())

# 1) CSV que abre certinho no Excel brasileiro
df.to_csv(
    SAIDA / "medicoes_status.csv", sep=";", decimal=",", index=False, encoding="utf-8-sig"
)

# 2) Excel com várias abas: um ExcelWriter, vários to_excel
resumo = df.groupby("linha", as_index=False).agg(
    bobinas=("bobina", "count"), peso_total_kg=("peso_kg", "sum"), cof_medio=("cof", "mean")
)
caminho = SAIDA / "relatorio_cof.xlsx"
with pd.ExcelWriter(caminho, engine="openpyxl") as escritor:
    resumo.to_excel(escritor, sheet_name="Resumo", index=False)
    df.to_excel(escritor, sheet_name="Medições", index=False)

    # 3) Formatar com openpyxl ANTES de sair do with (o arquivo é salvo no fim do bloco)
    cabecalho_fonte = Font(bold=True, color="FFFFFF")
    cabecalho_fundo = PatternFill("solid", fgColor="1F4E78")
    vermelho = PatternFill("solid", fgColor="FFC7CE")
    amarelo = PatternFill("solid", fgColor="FFEB9C")

    for nome_aba, tabela in [("Resumo", resumo), ("Medições", df)]:
        aba = escritor.sheets[nome_aba]
        for celula in aba[1]:  # linha 1 = cabeçalho
            celula.font = cabecalho_fonte
            celula.fill = cabecalho_fundo
        for i, coluna in enumerate(tabela.columns, start=1):
            maior = max(len(str(coluna)), *(len(str(v)) for v in tabela[coluna]))
            aba.column_dimensions[get_column_letter(i)].width = min(maior + 2, 40)
        aba.freeze_panes = "A2"  # congela o cabeçalho
        aba.auto_filter.ref = aba.dimensions  # liga o filtro do Excel

    aba = escritor.sheets["Resumo"]
    for linha in aba.iter_rows(min_row=2, min_col=3, max_col=4):
        linha[0].number_format = "#,##0.0"  # peso (o Excel mostra no padrão do Windows)
        linha[1].number_format = "0.000"  # COF médio

    aba = escritor.sheets["Medições"]
    col_status = df.columns.get_loc("status") + 1  # openpyxl conta colunas a partir de 1
    for linha in aba.iter_rows(min_row=2):
        status = linha[col_status - 1].value
        if status in ("REPROVADA", "LIBERADA"):
            fundo = vermelho if status == "REPROVADA" else amarelo
            for celula in linha:
                celula.fill = fundo

print("Gerados: saida/medicoes_status.csv e saida/relatorio_cof.xlsx")
