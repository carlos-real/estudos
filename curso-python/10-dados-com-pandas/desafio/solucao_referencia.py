"""Desafio do módulo 10 — Relatório mensal da qualidade (solução de referência).

Lê as medições, as OPs do ERP e as especificações, confere cada bobina do mês
escolhido e gera:
  - um Excel formatado com as abas Resumo, Não conformidades e Medições;
  - um gráfico de barras com o % de aprovação por linha;
  - um gráfico de controle (média ± 3σ) da espessura de um produto.

Uso (na pasta curso-python):
    python 10-dados-com-pandas/desafio/solucao_referencia.py
    python 10-dados-com-pandas/desafio/solucao_referencia.py --mes 2026-08
    python 10-dados-com-pandas/desafio/solucao_referencia.py --mes 2026-09 --produto PE-35

Sem --mes, usa o último mês que aparece nos dados.
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

import matplotlib

matplotlib.use("Agg")  # só arquivos, sem janelas

import matplotlib.pyplot as plt  # noqa: E402
import pandas as pd  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

PASTA_MODULO = Path(__file__).resolve().parent.parent
CORES = {"REPROVADA": "FFC7CE", "PENDENTE": "FFEB9C"}


# --------------------------------------------------------------------------- leitura
def ler_dados(pasta: Path) -> pd.DataFrame:
    """Medições + produto/cliente (ERP) + especificação, numa tabela só."""
    medicoes = pd.read_csv(
        pasta / "medicoes.csv", sep=";", decimal=",", thousands=".", dtype={"op": str}
    )
    ops = pd.read_csv(
        pasta / "ops_protheus.csv",
        sep=";",
        decimal=",",
        thousands=".",
        encoding="cp1252",
        dtype={"op": str},
    )
    specs = pd.read_csv(pasta / "especificacoes.csv", sep=";", decimal=",")

    medicoes["data"] = pd.to_datetime(medicoes["data"], format="%d/%m/%Y")
    medicoes["mes"] = medicoes["data"].dt.strftime("%Y-%m")
    df = medicoes.merge(ops[["op", "cliente"]], on="op", how="left", validate="many_to_one")
    df = df.merge(specs, on="produto", how="left", validate="many_to_one")
    return df


# --------------------------------------------------------------------------- regras
def conferir(df: pd.DataFrame) -> pd.DataFrame:
    """Acrescenta 'status' (APROVADA, REPROVADA, PENDENTE) e 'motivo'."""
    df = df.copy()
    tol = df["tolerancia_pct"] / 100
    esp_fora = ~df["espessura_um"].between(
        df["espessura_nominal"] * (1 - tol), df["espessura_nominal"] * (1 + tol)
    )
    gram_fora = ~df["gramatura_gm2"].between(
        df["gramatura_nominal"] * (1 - tol), df["gramatura_nominal"] * (1 + tol)
    )
    cof_fora = df["cof"] > df["cof_max"]
    liberado = df["observacao"].fillna("").str.contains("cof: liberado", case=False)

    df["motivo"] = ""
    df.loc[esp_fora, "motivo"] += "espessura; "
    df.loc[gram_fora, "motivo"] += "gramatura; "
    df.loc[cof_fora & ~liberado, "motivo"] += "COF; "
    df["motivo"] = df["motivo"].str.rstrip("; ")

    df["status"] = "APROVADA"
    df.loc[df["motivo"] != "", "status"] = "REPROVADA"
    # Pendências por último: elas vencem as regras acima
    sem_spec = df["espessura_nominal"].isna()
    df.loc[sem_spec, ["status", "motivo"]] = ["PENDENTE", "produto sem especificação"]
    sem_erp = df["cliente"].isna()
    df.loc[sem_erp, ["status", "motivo"]] = ["PENDENTE", "OP sem cadastro no ERP"]
    return df


def resumir(df: pd.DataFrame) -> pd.DataFrame:
    """Uma linha por linha de produção, mais a linha de total."""
    contagem = pd.crosstab(df["linha"], df["status"])
    for status in ["APROVADA", "REPROVADA", "PENDENTE"]:
        if status not in contagem.columns:
            contagem[status] = 0
    resumo = contagem[["APROVADA", "REPROVADA", "PENDENTE"]].copy()
    resumo.columns = ["aprovadas", "reprovadas", "pendentes"]
    resumo["peso_total_kg"] = df.groupby("linha")["peso_kg"].sum().round(1)
    resumo.loc["Total"] = resumo.sum()
    # A linha de total mistura inteiros e o peso (float): o pandas converte tudo para
    # float. Voltamos as contagens para inteiro.
    contagens = ["aprovadas", "reprovadas", "pendentes"]
    resumo[contagens] = resumo[contagens].astype(int)
    resumo.insert(0, "bobinas", resumo[contagens].sum(axis=1))
    resumo["pct_aprovadas"] = (resumo["aprovadas"] / resumo["bobinas"] * 100).round(1)
    return resumo.reset_index()


# --------------------------------------------------------------------------- saídas
def formatar_aba(aba, tabela: pd.DataFrame) -> None:
    """Cabeçalho, larguras, painel congelado e filtro."""
    for celula in aba[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="1F4E78")
    for i, coluna in enumerate(tabela.columns, start=1):
        maior = max([len(str(coluna)), *(len(str(v)) for v in tabela[coluna])])
        aba.column_dimensions[get_column_letter(i)].width = min(maior + 2, 45)
    aba.freeze_panes = "A2"
    aba.auto_filter.ref = aba.dimensions


def gravar_excel(caminho: Path, resumo: pd.DataFrame, mes_df: pd.DataFrame) -> None:
    colunas = ["data", "linha", "op", "cliente", "produto", "bobina", "espessura_um",
               "gramatura_gm2", "cof", "peso_kg", "observacao", "status", "motivo"]  # fmt: skip
    medicoes = mes_df[colunas].copy()
    medicoes["data"] = medicoes["data"].dt.date  # sem a hora 00:00:00 no Excel
    nao_conformes = medicoes[medicoes["status"] != "APROVADA"]

    with pd.ExcelWriter(caminho, engine="openpyxl") as escritor:
        for nome, tabela in [
            ("Resumo", resumo),
            ("Não conformidades", nao_conformes),
            ("Medições", medicoes),
        ]:
            tabela.to_excel(escritor, sheet_name=nome, index=False)
            formatar_aba(escritor.sheets[nome], tabela)

        for nome, tabela in [("Não conformidades", nao_conformes), ("Medições", medicoes)]:
            aba = escritor.sheets[nome]
            posicao = list(tabela.columns).index("status")
            for linha in aba.iter_rows(min_row=2):
                cor = CORES.get(linha[posicao].value)
                if cor:
                    for celula in linha:
                        celula.fill = PatternFill("solid", fgColor=cor)
            for linha in aba.iter_rows(min_row=2):  # datas no formato brasileiro
                linha[0].number_format = "DD/MM/YYYY"

        aba = escritor.sheets["Resumo"]
        aba.cell(row=aba.max_row, column=1).font = Font(bold=True)  # linha do Total


def grafico_aprovacao(resumo: pd.DataFrame, mes: str, caminho: Path) -> None:
    por_linha = resumo[resumo["linha"] != "Total"].set_index("linha")["pct_aprovadas"]
    fig, ax = plt.subplots(figsize=(6, 3.5))
    por_linha.plot.bar(ax=ax, color="#1F4E78", rot=0)
    ax.set_ylim(0, 105)
    ax.set_title(f"Bobinas aprovadas por linha ({mes})")
    ax.set_xlabel("Linha")
    ax.set_ylabel("Aprovadas (%)")
    for x, valor in enumerate(por_linha):
        ax.text(x, valor + 1, f"{valor:.0f}%", ha="center", fontsize=9)
    fig.tight_layout()
    fig.savefig(caminho, dpi=120)
    plt.close(fig)


def grafico_controle(df: pd.DataFrame, produto: str, mes: str, caminho: Path) -> str:
    """Limites calculados nos meses ANTERIORES (referência); vigia o mês escolhido."""
    do_produto = df[df["produto"] == produto].sort_values(["data", "bobina"])
    referencia = do_produto[do_produto["mes"] < mes]["espessura_um"]
    vigiados = do_produto[do_produto["mes"] == mes]["espessura_um"].reset_index(drop=True)
    if vigiados.empty:
        return f"sem bobinas de {produto} em {mes}: gráfico de controle não gerado"
    aviso = ""
    if len(referencia) < 2:
        referencia = vigiados
        aviso = " (sem histórico: limites calculados no próprio mês)"
    media, sigma = referencia.mean(), referencia.std()
    lic, lsc = media - 3 * sigma, media + 3 * sigma
    fora = vigiados[(vigiados < lic) | (vigiados > lsc)]

    fig, ax = plt.subplots(figsize=(7, 3.5))
    ax.plot(range(1, len(vigiados) + 1), vigiados, marker="o", label="Espessura")
    ax.axhline(media, color="green", label=f"Média {media:.1f}")
    ax.axhline(lsc, color="red", linestyle="--", label=f"LSC {lsc:.1f}")
    ax.axhline(lic, color="red", linestyle="--", label=f"LIC {lic:.1f}")
    if not fora.empty:
        ax.scatter(fora.index + 1, fora, color="red", s=80, zorder=3)
    ax.set_title(f"Controle da espessura — {produto} ({mes})")
    ax.set_xlabel("Bobina (ordem de medição)")
    ax.set_ylabel("Espessura (µm)")
    ax.legend(fontsize=8)
    fig.tight_layout()
    fig.savefig(caminho, dpi=120)
    plt.close(fig)
    return (
        f"{produto}: média {media:.2f} µm, LIC {lic:.2f}, LSC {lsc:.2f}, "
        f"{len(fora)} ponto(s) fora de controle{aviso}"
    )


# --------------------------------------------------------------------------- CLI
def criar_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Relatório mensal da qualidade.")
    parser.add_argument("--mes", help="mês no formato AAAA-MM (padrão: o último dos dados)")
    parser.add_argument("--produto", default="LAM-62", help="produto do gráfico de controle")
    parser.add_argument("--dados", type=Path, default=PASTA_MODULO / "dados")
    parser.add_argument("--saida", type=Path, default=PASTA_MODULO / "saida")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = criar_parser().parse_args(argv)
    try:
        df = conferir(ler_dados(args.dados))
    except FileNotFoundError as erro:
        print(f"Erro: arquivo não encontrado: {erro.filename}", file=sys.stderr)
        return 1

    mes = args.mes or df["mes"].max()
    mes_df = df[df["mes"] == mes]
    if mes_df.empty:
        meses = ", ".join(sorted(df["mes"].unique()))
        print(f"Erro: não há medições em {mes}. Meses disponíveis: {meses}", file=sys.stderr)
        return 1

    args.saida.mkdir(parents=True, exist_ok=True)
    resumo = resumir(mes_df)
    excel = args.saida / f"relatorio_qualidade_{mes}.xlsx"
    gravar_excel(excel, resumo, mes_df)
    grafico_aprovacao(resumo, mes, args.saida / f"aprovacao_{mes}.png")
    info_controle = grafico_controle(df, args.produto, mes, args.saida / f"controle_{mes}.png")

    print(f"Relatório de {mes}: {len(mes_df)} bobinas")
    print(resumo.to_string(index=False))
    print(info_controle)
    print(f"Arquivos em {args.saida}: {excel.name}, aprovacao_{mes}.png, controle_{mes}.png")
    return 0


if __name__ == "__main__":
    sys.exit(main())
