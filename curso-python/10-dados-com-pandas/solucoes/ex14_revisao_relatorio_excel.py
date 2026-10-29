"""
Exercício 10.14 — Revisão: relatório Excel formatado a partir de objetos (SOLUÇÃO)

Fluxo: objetos (módulo 08) -> DataFrame -> Excel (pandas) -> formatação (openpyxl).
O pandas grava os dados; o openpyxl, que é quem o pandas usa por baixo, deixa a
planilha bonita. As duas coisas acontecem dentro do mesmo "with ExcelWriter".
"""

from dataclasses import asdict, dataclass, fields
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

VERMELHO_CLARO = PatternFill("solid", fgColor="FFC7CE")


@dataclass
class Resultado:
    """Resultado da conferência de uma bobina (já pronto, não precisa mexer)."""

    op: str
    bobina: str
    espessura_um: float
    status: str


def resultados_para_dataframe(resultados: list[Resultado]) -> pd.DataFrame:
    """Lista de Resultado -> DataFrame com as colunas na ordem dos campos."""
    # fields() lista os campos da dataclass na ordem da declaração. Usar isso em
    # columns=... garante a ordem E as colunas mesmo com a lista vazia.
    colunas = [campo.name for campo in fields(Resultado)]
    return pd.DataFrame([asdict(r) for r in resultados], columns=colunas)


def exportar_relatorio(resultados: list[Resultado], caminho: str | Path) -> Path:
    """Grava o relatório formatado e devolve o caminho do arquivo."""
    if not resultados:
        raise ValueError("não há resultados para exportar")

    caminho = Path(caminho)
    caminho.parent.mkdir(parents=True, exist_ok=True)
    df = resultados_para_dataframe(resultados)

    with pd.ExcelWriter(caminho, engine="openpyxl") as escritor:
        df.to_excel(escritor, sheet_name="Resultados", index=False)
        aba = escritor.sheets["Resultados"]  # a planilha do openpyxl

        for celula in aba[1]:  # aba[1] = linha 1 (o cabeçalho)
            celula.font = Font(bold=True)
        aba.freeze_panes = "A2"

        # Largura: o maior entre o cabeçalho e os valores, com folga de 2.
        for numero, coluna in enumerate(df.columns, start=1):
            maior = max([len(coluna), *(len(str(v)) for v in df[coluna])])
            aba.column_dimensions[get_column_letter(numero)].width = maior + 2

        # Linha i do DataFrame é a linha i + 2 da planilha (1 é o cabeçalho e o
        # openpyxl começa a contar do 1).
        for i, status in enumerate(df["status"]):
            if status == "REPROVADA":
                for celula in aba[i + 2]:
                    celula.fill = VERMELHO_CLARO
    # Ao sair do with, o arquivo é salvo (com a formatação).
    return caminho
