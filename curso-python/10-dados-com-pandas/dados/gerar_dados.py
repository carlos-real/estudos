"""Gera os arquivos de dados usados na aula do módulo 10.

Os dados são FICTÍCIOS, mas imitam o que aparece num laboratório de qualidade de
embalagens flexíveis: medições de bobinas, OPs exportadas do ERP, especificações
de produto e a planilha de valores com uma aba por linha de produção.

Uso (na pasta curso-python):
    python 10-dados-com-pandas/dados/gerar_dados.py

O resultado é sempre o mesmo (a semente do sorteio é fixa). Se apagar algum
arquivo da pasta dados/, rode de novo para recriá-lo.
"""

from __future__ import annotations

import csv
import random
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

PASTA = Path(__file__).resolve().parent
ARQUIVO_SALARIOS = PASTA.parent.parent.parent / "arquivos" / "salaries.csv"

# produto: (descrição, linha, espessura nominal µm, tolerância %, gramatura nominal, COF máx.)
# As linhas seguem o projeto final: A e B extrusam PE, C coextrusa, D lamina.
PRODUTOS = {
    "PE-40": ("Filme PE 40 µm transparente", "A", 40.0, 10.0, 37.0, 0.35),
    "PE-50": ("Filme PE 50 µm branco", "A", 50.0, 10.0, 46.3, 0.35),
    "PE-35": ("Filme PE 35 µm", "B", 35.0, 10.0, 32.4, 0.40),
    "COEX-70": ("Filme coextrusado 70 µm", "C", 70.0, 8.0, 65.0, 0.30),
    "LAM-62": ("Laminado PET/PE 62 µm", "D", 62.0, 7.0, 62.7, 0.30),
}

CLIENTES = [
    "Torrefação Serra Ltda",
    "Agro Campo S.A.",
    "Biscoitos Aurora",
    "Laticínios Serra Azul",
    "Higiene Bem-Estar S.A.",
]

# (op, produto, data de emissão, índice do cliente)
# No Protheus o número da OP tem 6 caracteres, com zeros à esquerda.
OPS = [
    ("024412", "PE-40", date(2026, 7, 2), 0),
    ("024413", "PE-35", date(2026, 7, 3), 3),
    ("024420", "COEX-70", date(2026, 7, 8), 1),
    ("024421", "LAM-62", date(2026, 7, 10), 0),
    ("024430", "PE-50", date(2026, 7, 15), 2),
    ("024433", "PE-35", date(2026, 7, 21), 4),
    ("024440", "PE-40", date(2026, 7, 28), 2),
    ("024447", "LAM-62", date(2026, 8, 4), 0),
    ("024451", "COEX-70", date(2026, 8, 6), 1),
    ("024458", "PE-35", date(2026, 8, 12), 3),
    ("024462", "PE-40", date(2026, 8, 18), 4),
    ("024466", "LAM-62", date(2026, 8, 25), 2),
    ("024470", "PE-50", date(2026, 9, 1), 1),
    ("024474", "COEX-70", date(2026, 9, 3), 3),
    ("024479", "PE-35", date(2026, 9, 9), 4),
    ("024483", "PE-90", date(2026, 9, 15), 1),  # produto novo: ainda sem especificação
    ("024488", "LAM-62", date(2026, 9, 22), 0),
]
OP_SEM_CADASTRO = "024499"  # aparece nas medições, mas não no ERP (para o merge)

# Casos planejados, para a aula ter exemplos de COF fora (liberado ou não)
AJUSTES = {
    "024420-03": {"cof": 0.34, "observacao": "COF: liberado"},
    "024458-02": {"cof": 0.44},
}


def br(valor: float, casas: int = 1) -> str:
    """Formata um número no padrão brasileiro: 1250.5 -> '1.250,5'."""
    texto = f"{valor:,.{casas}f}"
    return texto.replace(",", "X").replace(".", ",").replace("X", ".")


def gerar_medicoes(sorteio: random.Random) -> list[dict]:
    medicoes = []
    for op, produto, emissao, _ in [*OPS, (OP_SEM_CADASTRO, "PE-40", date(2026, 9, 25), 0)]:
        linha = PRODUTOS[produto][1] if produto in PRODUTOS else "A"
        nominal = PRODUTOS[produto][2] if produto in PRODUTOS else 90.0
        gram_nominal = PRODUTOS[produto][4] if produto in PRODUTOS else 82.8
        cof_max = PRODUTOS[produto][5] if produto in PRODUTOS else 0.35
        for numero in range(1, sorteio.randint(3, 5) + 1):
            espessura = round(sorteio.gauss(nominal, nominal * 0.045), 1)
            gramatura = round(gram_nominal * espessura / nominal + sorteio.gauss(0, 0.4), 1)
            cof = round(sorteio.gauss(cof_max - 0.08, 0.04), 2)
            observacao = "COF: liberado" if cof > cof_max and sorteio.random() < 0.5 else ""
            medicoes.append(
                {
                    "data": emissao + timedelta(days=sorteio.randint(1, 3)),
                    "op": op,
                    "linha": linha,
                    "produto": produto,
                    "bobina": f"{op}-{numero:02d}",
                    "espessura_um": espessura,
                    "gramatura_gm2": gramatura,
                    "cof": cof,
                    "peso_kg": round(sorteio.uniform(780, 1650), 1),
                    "observacao": observacao,
                }
            )
            medicoes[-1].update(AJUSTES.get(medicoes[-1]["bobina"], {}))
    return medicoes


def gravar_medicoes_csv(medicoes: list[dict]) -> None:
    """CSV no padrão do Excel brasileiro: ponto e vírgula, vírgula decimal, UTF-8 com BOM."""
    caminho = PASTA / "medicoes.csv"
    with caminho.open("w", encoding="utf-8-sig", newline="") as arquivo:
        escritor = csv.writer(arquivo, delimiter=";")
        escritor.writerow(list(medicoes[0]))
        for m in medicoes:
            escritor.writerow(
                [
                    m["data"].strftime("%d/%m/%Y"),
                    m["op"],
                    m["linha"],
                    m["produto"],
                    m["bobina"],
                    br(m["espessura_um"]),
                    br(m["gramatura_gm2"]),
                    br(m["cof"], 2),
                    br(m["peso_kg"]),
                    m["observacao"],
                ]
            )


def gravar_medicoes_sujas_csv() -> None:
    """Um arquivo pequeno com os defeitos típicos de planilha digitada à mão."""
    linhas = [
        ["data", "op", "linha", "bobina", "espessura_um", "peso_kg", "operador"],
        ["01/09/2026", "024470", "A", "024470-01", "51,2", "1.120,5", "Marcos"],
        ["01/09/2026", "024470", " a", "024470-02", "49,8", "1.098,0", "marcos "],
        ["01/09/2026", "024470", "A", "024470-02", "49,8", "1.098,0", "marcos "],
        ["02/09/2026", "024470", "A ", "024470-03", "n/d", "1.205,3", "Marcos"],
        ["04/09/2026", "024474", "c", "024474-01", "71,3", "845,0", "Júlia"],
        ["04/09/2026", "024474", "C", "024474-02", "", "910,2", "JÚLIA"],
        ["04/09/2026", "024474", "C", "024474-03", "69,9", "", "Júlia"],
        ["", "024474", "C", "024474-04", "70,1", "877,7", "Júlia"],
        ["10/09/2026", "024479", "b", "024479-01", "35,4", "1.310,0", "Ana"],
        ["10/09/2026", "024479", "B", "024479-02", "34,7 ", "1.288,4", "ana"],
    ]
    caminho = PASTA / "medicoes_sujas.csv"
    with caminho.open("w", encoding="utf-8-sig", newline="") as arquivo:
        csv.writer(arquivo, delimiter=";").writerows(linhas)


def gravar_ops_protheus_csv(sorteio: random.Random) -> None:
    """Exportação do ERP: separador ';' e encoding cp1252 (padrão antigo do Windows)."""
    caminho = PASTA / "ops_protheus.csv"
    with caminho.open("w", encoding="cp1252", newline="") as arquivo:
        escritor = csv.writer(arquivo, delimiter=";")
        escritor.writerow(["op", "cliente", "produto", "emissao", "quantidade_kg"])
        for op, produto, emissao, cliente in OPS:
            quantidade = sorteio.randrange(2_000, 6_001, 250)
            escritor.writerow(
                [op, CLIENTES[cliente], produto, emissao.strftime("%d/%m/%Y"), br(quantidade)]
            )


def gravar_especificacoes_csv() -> None:
    caminho = PASTA / "especificacoes.csv"
    with caminho.open("w", encoding="utf-8-sig", newline="") as arquivo:
        escritor = csv.writer(arquivo, delimiter=";")
        cabecalho = ["produto", "descricao", "espessura_nominal", "tolerancia_pct"]
        escritor.writerow([*cabecalho, "gramatura_nominal", "cof_max"])
        for produto, (descricao, _, nominal, tol, gram, cof) in PRODUTOS.items():
            escritor.writerow([produto, descricao, br(nominal), br(tol), br(gram), br(cof, 2)])


def gravar_planilha_linhas(medicoes: list[dict]) -> None:
    """A 'planilha de valores' do laboratório: uma aba por linha de produção.

    Como muita planilha feita à mão, ela tem um título na linha 1 e o cabeçalho
    só na linha 2. Os nomes das colunas têm acento e unidade.
    """
    cabecalho = [
        "OP",
        "Bobina",
        "Data",
        "Espessura (µm)",
        "Gramatura (g/m²)",
        "COF",
        "Peso (kg)",
        "Observação",
    ]
    livro = Workbook()
    livro.remove(livro.active)
    for linha in "ABCD":
        aba = livro.create_sheet(f"Linha {linha}")
        aba["A1"] = f"Controle de qualidade — Linha {linha}"
        aba["A1"].font = Font(bold=True, size=13)
        aba.append(cabecalho)
        for celula in aba[2]:
            celula.font = Font(bold=True, color="FFFFFF")
            celula.fill = PatternFill("solid", fgColor="1F4E78")
            celula.alignment = Alignment(horizontal="center")
        for m in medicoes:
            if m["linha"] != linha:
                continue
            aba.append(
                [
                    m["op"],
                    m["bobina"],
                    m["data"],
                    m["espessura_um"],
                    m["gramatura_gm2"],
                    m["cof"],
                    m["peso_kg"],
                    m["observacao"] or None,
                ]
            )
        for linha_planilha in aba.iter_rows(min_row=3, min_col=3, max_col=3):
            linha_planilha[0].number_format = "DD/MM/YYYY"
        for coluna, largura in zip("ABCDEFGH", [10, 12, 12, 15, 17, 8, 11, 16], strict=True):
            aba.column_dimensions[coluna].width = largura
        aba.freeze_panes = "A3"
    livro.save(PASTA / "planilha_linhas.xlsx")


def gravar_amostra_salarios() -> None:
    """Amostra pequena do arquivo de salários, para estudar sem o arquivo grande."""
    if not ARQUIVO_SALARIOS.exists():
        print(f"(pulei a amostra de salários: {ARQUIVO_SALARIOS} não existe)")
        return
    import pandas as pd

    df = pd.read_csv(ARQUIVO_SALARIOS)
    amostra = df.sample(n=600, random_state=42).sort_index()
    amostra.to_csv(PASTA / "salarios_amostra.csv", index=False)


def main() -> None:
    sorteio = random.Random(2026)
    medicoes = gerar_medicoes(sorteio)
    gravar_medicoes_csv(medicoes)
    gravar_medicoes_sujas_csv()
    gravar_ops_protheus_csv(sorteio)
    gravar_especificacoes_csv()
    gravar_planilha_linhas(medicoes)
    gravar_amostra_salarios()
    for arquivo in sorted(PASTA.iterdir()):
        if arquivo.suffix in {".csv", ".xlsx"}:
            print(f"{arquivo.name:<24} {arquivo.stat().st_size / 1024:6.1f} KB")


if __name__ == "__main__":
    main()
