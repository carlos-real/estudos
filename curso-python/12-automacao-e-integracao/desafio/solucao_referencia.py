"""Desafio 12 — Robô da expedição (solução de referência).

Lê os pedidos de certificado salvos como .eml, confere cada OP no cadastro,
monta a fila de emissão, "emite" (com um emissor simulado no lugar do robô
Playwright), gera o relatório Excel e o e-mail-resumo (sem enviar).

Uso (na pasta curso-python):
    python 12-automacao-e-integracao/desafio/solucao_referencia.py
    python 12-automacao-e-integracao/desafio/solucao_referencia.py --confirmar
    python 12-automacao-e-integracao/desafio/solucao_referencia.py --emails PASTA --ops ops.csv

Por padrão roda em MODO ENSAIO: planeja tudo, gera relatório, mas não registra
nenhuma emissão. Rodar duas vezes com --confirmar não emite nada em dobro.
"""

from __future__ import annotations

import argparse
import csv
import logging
import re
import sys
from dataclasses import dataclass
from datetime import date
from email import policy
from email.message import EmailMessage
from email.parser import BytesParser
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

PASTA_MODULO = Path(__file__).resolve().parent.parent
PADRAO_EMAILS = PASTA_MODULO / "dados" / "emails"
PADRAO_OPS = PASTA_MODULO / "dados" / "ops.csv"
PADRAO_SAIDA = Path(__file__).resolve().parent / "saida"

RE_OP = re.compile(r"\bOP\s*:?\s*(\d{6})\b", re.IGNORECASE)
RE_NF = re.compile(r"\bNF\s*:?\s*(\d+)", re.IGNORECASE)
RE_PESO = re.compile(r"\bPeso\s*:?\s*([\d.]+(?:,\d+)?)\s*kg", re.IGNORECASE)
RE_TAG = re.compile(r"<[^>]+>")

log = logging.getLogger("robo_expedicao")


# ----------------------------------------------------------------------------
# Modelo de dados
# ----------------------------------------------------------------------------
@dataclass
class Pedido:
    arquivo: str
    op: str
    nf: str
    peso_kg: float
    linha: str = ""
    produto: str = ""
    especificacao: str = ""
    status: str = "NA FILA"


# ----------------------------------------------------------------------------
# Leitura: e-mails e cadastro de OPs
# ----------------------------------------------------------------------------
def numero_br(texto: str) -> float:
    return float(texto.replace(".", "").replace(",", "."))


def formatar_br(valor: float) -> str:
    return f"{valor:,.2f}".replace(",", "#").replace(".", ",").replace("#", ".")


def ler_email(caminho: Path) -> Pedido:
    """Extrai OP, NF e peso. Levanta ValueError dizendo o que faltou."""
    msg = BytesParser(policy=policy.default).parsebytes(caminho.read_bytes())
    corpo = msg.get_body(preferencelist=("plain", "html"))
    texto = corpo.get_content() if corpo is not None else ""
    if corpo is not None and corpo.get_content_type() == "text/html":
        texto = RE_TAG.sub(" ", texto)
    achados = {"OP": RE_OP.search(texto), "NF": RE_NF.search(texto), "peso": RE_PESO.search(texto)}
    for campo, achado in achados.items():
        if achado is None:
            raise ValueError(f"{campo} não encontrado no e-mail")
    return Pedido(
        arquivo=caminho.name,
        op=achados["OP"].group(1),
        nf=achados["NF"].group(1),
        peso_kg=numero_br(achados["peso"].group(1)),
    )


def carregar_ops(caminho: Path) -> dict[str, dict]:
    with caminho.open(encoding="utf-8", newline="") as arquivo:
        return {linha["op"]: linha for linha in csv.DictReader(arquivo, delimiter=";")}


def montar_fila(pasta: Path, ops: dict[str, dict]) -> tuple[list[Pedido], list[tuple[str, str]]]:
    """Devolve (fila de pedidos válidos, pendências [(arquivo, motivo)])."""
    fila: list[Pedido] = []
    pendencias: list[tuple[str, str]] = []
    for caminho in sorted(pasta.glob("*.eml")):
        try:
            pedido = ler_email(caminho)
        except ValueError as erro:
            pendencias.append((caminho.name, str(erro)))
            continue
        cadastro = ops.get(pedido.op)
        if cadastro is None:
            pendencias.append((caminho.name, f"OP {pedido.op} sem cadastro"))
            continue
        pedido.linha = cadastro["linha"]
        pedido.produto = cadastro["produto"]
        titulo = cadastro["especificacao_titulo"].strip()
        if titulo:
            pedido.especificacao = f"{titulo}: {cadastro['especificacao_texto']}"
        else:
            pedido.especificacao = f"Padrão da linha {pedido.linha}"
        fila.append(pedido)
    return fila, pendencias


# ----------------------------------------------------------------------------
# Robô: plano, emissor (simulado) e registro idempotente
# ----------------------------------------------------------------------------
def planejar(pedido: Pedido, ensaio: bool) -> list[tuple[str, str, str]]:
    acoes = [
        ("preencher", "OP", pedido.op),
        ("preencher", "NF", pedido.nf),
        ("preencher", "Peso (kg)", formatar_br(pedido.peso_kg)),
        ("selecionar", "Linha", pedido.linha),
        ("preencher", "Observação", pedido.especificacao),
        ("print", f"antes_de_confirmar_{pedido.op}", ""),
        ("clicar", "Cancelar" if ensaio else "Confirmar", ""),
    ]
    return acoes


class EmissorSimulado:
    """Faz o papel do robô Playwright: devolve números de certificado em sequência.

    Para usar o robô de verdade, troque por uma função que execute o plano no
    navegador (veja exemplos/09_robo_certificados.py) e devolva o número.
    """

    def __init__(self, inicio: int = 1):
        self.proximo = inicio

    def __call__(self, pedido: Pedido) -> str:
        numero = f"{self.proximo:06d}"
        self.proximo += 1
        return numero


def ler_registro(caminho: Path) -> set[tuple[str, str]]:
    if not caminho.exists():
        return set()
    with caminho.open(encoding="utf-8", newline="") as arquivo:
        return {(lin["op"], lin["nf"]) for lin in csv.DictReader(arquivo, delimiter=";")}


def anotar(caminho: Path, pedido: Pedido, certificado: str) -> None:
    novo = not caminho.exists()
    with caminho.open("a", encoding="utf-8", newline="") as arquivo:
        escritor = csv.writer(arquivo, delimiter=";")
        if novo:
            escritor.writerow(["op", "nf", "certificado"])
        escritor.writerow([pedido.op, pedido.nf, certificado])


def processar(fila: list[Pedido], emitir, registro: Path, ensaio: bool) -> None:
    """Atualiza pedido.status de cada item da fila."""
    feitos = ler_registro(registro)
    for pedido in fila:
        if (pedido.op, pedido.nf) in feitos:
            pedido.status = "JÁ EMITIDO"
            log.info("OP %s / NF %s já emitida antes: pulando", pedido.op, pedido.nf)
            continue
        acoes = planejar(pedido, ensaio)
        log.info("OP %s: plano com %d ações, termina em %s", pedido.op, len(acoes), acoes[-1][1])
        if ensaio:
            pedido.status = "ENSAIO"
            continue
        try:
            certificado = emitir(pedido)
        except Exception:
            log.exception("Falha ao emitir OP %s", pedido.op)
            pedido.status = "FALHA"
            continue
        anotar(registro, pedido, certificado)
        feitos.add((pedido.op, pedido.nf))
        pedido.status = f"EMITIDO {certificado}"
        log.info("OP %s: certificado %s emitido", pedido.op, certificado)


# ----------------------------------------------------------------------------
# Saídas: Excel e e-mail-resumo
# ----------------------------------------------------------------------------
def estilizar(ws, ultima_coluna: int) -> None:
    for celula in ws[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", start_color="1F4E78")
    for indice, coluna in enumerate(ws.iter_cols(max_col=ultima_coluna), start=1):
        maior = max(len(str(c.value)) for c in coluna if c.value is not None)
        ws.column_dimensions[get_column_letter(indice)].width = min(maior + 2, 60)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def gerar_excel(fila: list[Pedido], pendencias: list[tuple[str, str]], caminho: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Fila"
    ws.append(["Arquivo", "OP", "NF", "Peso (kg)", "Linha", "Produto", "Especificação", "Status"])
    for p in fila:
        ws.append([p.arquivo, p.op, p.nf, p.peso_kg, p.linha, p.produto, p.especificacao, p.status])
    for (celula,) in ws.iter_rows(min_row=2, min_col=4, max_col=4):
        celula.number_format = "#,##0.00"
    estilizar(ws, 8)
    if fila:
        ws.conditional_formatting.add(
            f"H2:H{ws.max_row}",
            CellIsRule(
                operator="equal",
                formula=['"FALHA"'],
                font=Font(color="9C0006"),
                fill=PatternFill("solid", start_color="FFC7CE", end_color="FFC7CE"),
            ),
        )

    wp = wb.create_sheet("Pendências")
    wp.append(["Arquivo", "Motivo"])
    for arquivo, motivo in pendencias:
        wp.append([arquivo, motivo])
    estilizar(wp, 2)

    caminho.parent.mkdir(parents=True, exist_ok=True)
    wb.save(caminho)
    return caminho


def montar_resumo(
    fila: list[Pedido], pendencias: list[tuple[str, str]], anexo: Path, dia: date, ensaio: bool
) -> EmailMessage:
    peso_total = sum(p.peso_kg for p in fila)
    linhas = [
        "Olá,",
        "",
        f"Resumo da fila de certificados de {dia:%d/%m/%Y}"
        + (" (MODO ENSAIO: nada foi emitido)." if ensaio else "."),
        "",
        f"Pedidos na fila: {len(fila)} ({formatar_br(peso_total)} kg)",
    ]
    linhas += [f"  - OP {p.op} / NF {p.nf}: {p.status}" for p in fila]
    linhas += ["", f"Pendências: {len(pendencias)}"]
    linhas += [f"  - {arquivo}: {motivo}" for arquivo, motivo in pendencias]
    linhas += ["", "Planilha completa em anexo.", "", "Robô da expedição"]

    msg = EmailMessage()
    msg["From"] = "robo.laboratorio@embalagens.example"
    msg["To"] = "supervisor.qualidade@embalagens.example"
    msg["Subject"] = f"Fila de certificados - {dia:%d/%m/%Y}"
    msg.set_content("\n".join(linhas) + "\n")
    msg.add_attachment(
        anexo.read_bytes(),
        maintype="application",
        subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=anexo.name,
    )
    return msg


# ----------------------------------------------------------------------------
# Programa principal
# ----------------------------------------------------------------------------
def configurar_log(saida: Path) -> None:
    saida.mkdir(parents=True, exist_ok=True)
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s %(levelname)-7s %(message)s",
        datefmt="%H:%M:%S",
        handlers=[
            logging.StreamHandler(sys.stdout),
            logging.FileHandler(saida / "robo_expedicao.log", encoding="utf-8"),
        ],
    )


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Robô da expedição (desafio 12).")
    parser.add_argument("--emails", type=Path, default=PADRAO_EMAILS, help="pasta com .eml")
    parser.add_argument("--ops", type=Path, default=PADRAO_OPS, help="cadastro de OPs (CSV ;)")
    parser.add_argument("--saida", type=Path, default=PADRAO_SAIDA, help="pasta de saída")
    parser.add_argument("--confirmar", action="store_true", help="registra as emissões")
    args = parser.parse_args(argv)
    ensaio = not args.confirmar

    configurar_log(args.saida)
    hoje = date.today()
    log.info("Início (%s)", "ENSAIO" if ensaio else "CONFIRMAR")

    if not args.emails.is_dir():
        log.error("Pasta de e-mails não encontrada: %s", args.emails)
        return 2

    fila, pendencias = montar_fila(args.emails, carregar_ops(args.ops))
    log.info("%d pedido(s) na fila, %d pendência(s)", len(fila), len(pendencias))
    for arquivo, motivo in pendencias:
        log.warning("Pendência em %s: %s", arquivo, motivo)

    processar(fila, EmissorSimulado(), args.saida / "emitidos.csv", ensaio)

    planilha = gerar_excel(fila, pendencias, args.saida / f"fila_{hoje:%Y-%m-%d}.xlsx")
    resumo = montar_resumo(fila, pendencias, planilha, hoje, ensaio)
    arquivo_eml = args.saida / f"resumo_{hoje:%Y-%m-%d}.eml"
    arquivo_eml.write_bytes(bytes(resumo))
    log.info("Gerados: %s e %s", planilha.name, arquivo_eml.name)

    print()
    print(resumo.get_body(preferencelist=("plain",)).get_content())
    return 1 if any(p.status == "FALHA" for p in fila) else 0


if __name__ == "__main__":
    sys.exit(main())
