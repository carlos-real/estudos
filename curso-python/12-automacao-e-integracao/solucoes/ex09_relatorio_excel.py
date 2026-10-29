"""
Exercício 12.09 — Relatório Excel de medições (SOLUÇÃO)

Estratégia: primeiro escrever TODOS os dados (append), depois aplicar o
estilo por cima. Misturar as duas coisas linha a linha deixa o código confuso.
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

CABECALHO = ["OP", "Linha", "Ensaio", "Valor", "Mínimo", "Máximo", "Status"]

# Estilos são objetos: criamos uma vez e reaproveitamos.
FONTE_CABECALHO = Font(bold=True, color="FFFFFF")
FUNDO_CABECALHO = PatternFill("solid", start_color="1F4E78")
FONTE_REPROVADO = Font(color="9C0006")
FUNDO_REPROVADO = PatternFill("solid", start_color="FFC7CE", end_color="FFC7CE")


def status(medicao: dict) -> str:
    """Regra de negócio isolada numa função pequena (fácil de testar e trocar)."""
    dentro = medicao["minimo"] <= medicao["valor"] <= medicao["maximo"]
    return "APROVADO" if dentro else "REPROVADO"


def gerar_relatorio(medicoes: list[dict], caminho: Path) -> Path:
    """Gera o relatório .xlsx das medições e devolve o caminho do arquivo."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Medições"

    # 1) Dados
    ws.append(CABECALHO)
    for m in medicoes:
        ws.append(
            [m["op"], m["linha"], m["ensaio"], m["valor"], m["minimo"], m["maximo"], status(m)]
        )

    # 2) Cabeçalho
    for celula in ws[1]:
        celula.font = FONTE_CABECALHO
        celula.fill = FUNDO_CABECALHO
        celula.alignment = Alignment(horizontal="center")

    # 3) Formato numérico só nas colunas D a F, a partir da linha 2.
    #    O "0.00" é o código do Excel; ele mostra "20,40" num Windows em português.
    for linha in ws.iter_rows(min_row=2, min_col=4, max_col=6):
        for celula in linha:
            celula.number_format = "0.00"

    # 4) Larguras (é o exercício 08 aplicado ao openpyxl)
    for indice, coluna in enumerate(ws.iter_cols(), start=1):
        maior = max(len(str(c.value)) for c in coluna if c.value is not None)
        ws.column_dimensions[get_column_letter(indice)].width = maior + 2

    # 5) Congelar cabeçalho e filtro. ws.max_row é a última linha com dados.
    ultima = ws.max_row
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:G{ultima}"

    # 6) Formatação condicional: quem pinta é o Excel, ao abrir o arquivo.
    #    Vantagem sobre pintar a célula direto: se alguém corrigir o valor
    #    para APROVADO, a cor some sozinha.
    if medicoes:
        ws.conditional_formatting.add(
            f"G2:G{ultima}",
            CellIsRule(
                operator="equal",
                formula=['"REPROVADO"'],  # aspas DENTRO do texto: é uma fórmula do Excel
                font=FONTE_REPROVADO,
                fill=FUNDO_REPROVADO,
            ),
        )

    caminho.parent.mkdir(parents=True, exist_ok=True)
    wb.save(caminho)
    return caminho
