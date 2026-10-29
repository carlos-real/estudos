"""Exemplo 05 — Relatório Excel com cara de profissional, usando openpyxl.

Gera saida/relatorio_medicoes.xlsx (ao lado deste script) com:
cabeçalho estilizado, larguras de coluna, formato numérico, painel congelado,
filtro e formatação condicional (REPROVADO em vermelho).

Rode com:  python exemplos/05_relatorio_excel.py
"""

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SAIDA = Path(__file__).resolve().parent / "saida"

MEDICOES = [
    {"op": "123456", "linha": "A", "ensaio": "Espessura", "valor": 20.4, "min": 18, "max": 22},
    {"op": "123456", "linha": "A", "ensaio": "COF", "valor": 0.27, "min": 0.2, "max": 0.3},
    {"op": "123789", "linha": "B", "ensaio": "Espessura", "valor": 43.1, "min": 38, "max": 42},
    {"op": "123789", "linha": "B", "ensaio": "Gramatura", "valor": 37.9, "min": 36.5, "max": 39.5},
    {"op": "124001", "linha": "C", "ensaio": "Tração", "valor": 21.5, "min": 25, "max": 40},
]

CABECALHO = ["OP", "Linha", "Ensaio", "Valor", "Mínimo", "Máximo", "Status"]


def status(medicao: dict) -> str:
    return "APROVADO" if medicao["min"] <= medicao["valor"] <= medicao["max"] else "REPROVADO"


def gerar(caminho: Path) -> None:
    wb = Workbook()
    ws = wb.active  # a planilha que já vem criada
    ws.title = "Medições"

    # 1) Dados: append() escreve uma linha inteira de uma vez
    ws.append(CABECALHO)
    for m in MEDICOES:
        ws.append([m["op"], m["linha"], m["ensaio"], m["valor"], m["min"], m["max"], status(m)])

    # 2) Estilo do cabeçalho
    for celula in ws[1]:  # ws[1] = linha 1 inteira
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", start_color="1F4E78")
        celula.alignment = Alignment(horizontal="center")

    # 3) Formato numérico (o Excel em português mostra 20,40)
    for linha in ws.iter_rows(min_row=2, min_col=4, max_col=6):
        for celula in linha:
            celula.number_format = "0.00"

    # 4) Largura: o maior texto da coluna + uma folga
    for indice, coluna in enumerate(ws.iter_cols(), start=1):
        maior = max(len(str(c.value)) for c in coluna if c.value is not None)
        ws.column_dimensions[get_column_letter(indice)].width = maior + 2

    # 5) Congelar o cabeçalho e ligar o filtro
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions  # ex.: "A1:G6"

    # 6) Formatação condicional: o EXCEL pinta, inclusive se alguém editar a célula
    ultima = ws.max_row
    ws.conditional_formatting.add(
        f"G2:G{ultima}",
        CellIsRule(
            operator="equal",
            formula=['"REPROVADO"'],
            font=Font(color="9C0006"),
            fill=PatternFill("solid", start_color="FFC7CE", end_color="FFC7CE"),
        ),
    )

    caminho.parent.mkdir(parents=True, exist_ok=True)
    wb.save(caminho)


def conferir(caminho: Path) -> None:
    """Reabre o arquivo e mostra o que foi gravado (é assim que os testes conferem)."""
    ws = load_workbook(caminho)["Medições"]
    print("Planilha:", ws.title, "| dimensões:", ws.dimensions)
    print("Cabeçalho negrito?", ws["A1"].font.bold, "| cor:", ws["A1"].fill.start_color.rgb)
    print("Largura da coluna C:", ws.column_dimensions["C"].width)
    print("Congelado em:", ws.freeze_panes, "| filtro:", ws.auto_filter.ref)
    for faixa in ws.conditional_formatting:
        regra = faixa.rules[0]
        print("Regra condicional:", faixa.sqref, regra.operator, regra.formula)
    reprovadas = [linha[0] for linha in ws.iter_rows(min_row=2, values_only=True)
                  if linha[6] == "REPROVADO"]
    print("OPs com reprovação:", reprovadas)


if __name__ == "__main__":
    arquivo = SAIDA / "relatorio_medicoes.xlsx"
    gerar(arquivo)
    print("Gerado:", arquivo.name)
    conferir(arquivo)
