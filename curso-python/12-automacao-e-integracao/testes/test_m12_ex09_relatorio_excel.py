from openpyxl import load_workbook

from _curso import carregar

ex = carregar(__file__, "ex09_relatorio_excel")

MEDICOES = [
    {
        "op": "123456",
        "linha": "A",
        "ensaio": "Espessura",
        "valor": 20.4,
        "minimo": 18.0,
        "maximo": 22.0,
    },
    {
        "op": "123456",
        "linha": "A",
        "ensaio": "COF",
        "valor": 0.31,
        "minimo": 0.2,
        "maximo": 0.3,
    },
    {
        "op": "123789",
        "linha": "B",
        "ensaio": "Tração longitudinal",
        "valor": 25.0,
        "minimo": 25.0,
        "maximo": 40.0,
    },
]


def gerar(tmp_path):
    """Gera o relatório e o reabre (é assim que conferimos o que foi gravado)."""
    caminho = tmp_path / "saida" / "turno.xlsx"
    devolvido = ex.gerar_relatorio(MEDICOES, caminho)
    assert devolvido == caminho, "A função deve devolver o caminho do arquivo"
    assert caminho.exists(), "O arquivo não foi criado (a pasta 'saida' existia?)"
    return load_workbook(caminho)["Medições"]


def test_cabecalho_e_dados(tmp_path):
    planilha = gerar(tmp_path)
    linhas = list(planilha.iter_rows(values_only=True))
    assert linhas[0] == ("OP", "Linha", "Ensaio", "Valor", "Mínimo", "Máximo", "Status")
    assert linhas[1] == ("123456", "A", "Espessura", 20.4, 18.0, 22.0, "APROVADO")
    assert len(linhas) == 4


def test_status_inclui_os_limites(tmp_path):
    planilha = gerar(tmp_path)
    status = [planilha.cell(row=linha, column=7).value for linha in range(2, 5)]
    assert status == ["APROVADO", "REPROVADO", "APROVADO"], (
        "Valor igual ao mínimo (ou ao máximo) está DENTRO da tolerância"
    )


def test_estilo_do_cabecalho(tmp_path):
    planilha = gerar(tmp_path)
    for celula in planilha[1]:
        assert celula.font.bold, f"{celula.coordinate} deveria estar em negrito"
        assert celula.font.color.rgb.endswith("FFFFFF")
        assert celula.fill.fill_type == "solid"
        assert celula.fill.start_color.rgb.endswith("1F4E78")


def test_formato_numerico(tmp_path):
    planilha = gerar(tmp_path)
    for coluna in "DEF":
        assert planilha[f"{coluna}2"].number_format == "0.00"
        assert planilha[f"{coluna}4"].number_format == "0.00"


def test_larguras_das_colunas(tmp_path):
    planilha = gerar(tmp_path)
    assert planilha.column_dimensions["C"].width == len("Tração longitudinal") + 2
    assert planilha.column_dimensions["A"].width == len("123456") + 2
    assert planilha.column_dimensions["G"].width == len("REPROVADO") + 2


def test_painel_congelado_e_filtro(tmp_path):
    planilha = gerar(tmp_path)
    assert planilha.freeze_panes == "A2"
    assert planilha.auto_filter.ref == "A1:G4"


def test_formatacao_condicional_no_status(tmp_path):
    planilha = gerar(tmp_path)
    regras = [
        (str(faixa.sqref), regra)
        for faixa in planilha.conditional_formatting
        for regra in faixa.rules
    ]
    assert regras, "Nenhuma formatação condicional encontrada"
    faixa, regra = regras[0]
    assert faixa == "G2:G4"
    assert regra.type == "cellIs"
    assert regra.operator == "equal"
    assert regra.formula == ['"REPROVADO"']


def test_sem_medicoes_gera_so_o_cabecalho(tmp_path):
    caminho = tmp_path / "vazio.xlsx"
    ex.gerar_relatorio([], caminho)
    ws = load_workbook(caminho)["Medições"]
    assert ws.max_row == 1
    assert ws["G1"].value == "Status"
    assert ws.auto_filter.ref == "A1:G1"
